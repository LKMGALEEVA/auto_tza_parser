"""
Модуль для сравнения данных из технологических схем и таблиц.

Содержит класс mTP для анализа и сопоставления кодов KKS из разных источников.
"""
import pandas as pd
import re
import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill


class mTP:
    """
    Класс для парсинга и сравнения технологических данных.
    
    Сравнивает коды KKS из схем с данными из табличных документов
    и формирует отчеты о расхождениях.
    """
    
    def parse_two_tables(self, scheme_file, table_file):
        """
        Сравнивает данные из файла схемы с файлом таблицы.
        
        Args:
            scheme_file: Путь к Excel файлу со схемой
            table_file: Путь к Excel файлу с таблицей
            
        Returns:
            list: Список результатов сравнения
        """
        df_scheme = pd.read_excel(scheme_file)
        df_table = pd.read_excel(table_file)

        # Извлечение данных из схемы
        scheme_values = df_scheme.iloc[:, 0]
        scheme_classes = (df_scheme.iloc[:, 1] if df_scheme.shape[1] > 1 
                         else pd.Series([""] * len(df_scheme)))
        
        # Извлечение данных из таблицы
        table_values = df_table.iloc[3:, 0]

        # Создание карты объектов из схемы
        objects_map = self._build_objects_map(scheme_values, scheme_classes)

        # Сравнение данных
        results = []
        used_keys = set()
        table_category = self._category_from_filename(table_file)

        # Поиск совпадений
        for raw_value in table_values:
            if pd.isna(raw_value):
                continue
                
            value = str(raw_value).strip()
            if not value or value.lower() == "nan":
                continue

            matched_keys = self._find_matching_keys(value, objects_map)

            if matched_keys:
                for key, category in matched_keys:
                    results.append(
                        self._build_result_row(key, "Совпадает", "green", category)
                    )
                    used_keys.add(key)
            else:
                results.append(
                    self._build_result_row(
                        value, "Есть в таблице, нет на схеме", "red", table_category
                    )
                )

        # Добавление объектов из схемы, не найденных в таблице
        for key, category in objects_map.items():
            if key not in used_keys:
                results.append(
                    self._build_result_row(
                        key, "Есть на схеме, нет в таблице", "red", category
                    )
                )

        return results
    
    def _build_objects_map(self, values, classes):
        """
        Создает словарь из данных схемы.
        
        Args:
            values: Значения из первого столбца
            classes: Классы объектов из второго столбца
            
        Returns:
            dict: Словарь {значение: категория}
        """
        objects_map = {}
        
        for raw_value, raw_class in zip(values, classes):
            if pd.isna(raw_value):
                continue
                
            value = str(raw_value).strip()
            if not value or value.lower() == "nan":
                continue
                
            category = self._category_from_raw(raw_class) or "Контроль"
            objects_map[value] = category
            
        return objects_map
    
    def _find_matching_keys(self, value, objects_map):
        """
        Находит соответствующие ключи для заданного значения.
        
        Args:
            value: Значение для поиска
            objects_map: Словарь объектов
            
        Returns:
            list: Список кортежей (ключ, категория)
        """
        matched_keys = []
        
        for key, category in objects_map.items():
            if re.match(r'^' + re.escape(key) + r'\b', value):
                matched_keys.append((key, category))
                
        return matched_keys

    def _write_results_to_sheet(self, ws, results):
        """
        Записывает результаты сравнения в лист Excel.
        
        Args:
            ws: Лист Excel для записи
            results: Список результатов для записи
        """
        # Заголовки
        headers = ["Значение", "Статус"]
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)

        # Данные
        for row_num, row in enumerate(results, 2):
            ws.cell(row=row_num, column=1, value=row["value"])
            ws.cell(row=row_num, column=2, value=row["status"])

        # Цветовое кодирование
        color_map = {
            "green": PatternFill(
                start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
            ),
            "red": PatternFill(
                start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"
            )
        }

        for row_num, row in enumerate(results, 2):
            fill_style = color_map.get(row["color"], color_map["red"])
            ws.cell(row=row_num, column=1).fill = fill_style
            ws.cell(row=row_num, column=2).fill = fill_style

        # Автоподгон ширины столбцов
        self._adjust_column_widths(ws)

    def _adjust_column_widths(self, ws):
        """
        Автоматически подгоняет ширину столбцов под содержимое.
        
        Args:
            ws: Лист Excel для обработки
        """
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            
            for cell in col:
                try:
                    cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length
                except:
                    pass
                    
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
    
    def _deduplicate_results(self, results):
        """
        Удаляет дубликаты из результатов.
        
        Args:
            results: Список результатов
            
        Returns:
            list: Список уникальных результатов
        """
        unique_results = []
        seen = set()
        
        for row in results:
            row_key = (row["value"], row["status"], row.get("category"))
            if row_key not in seen:
                seen.add(row_key)
                unique_results.append(row)
                
        return unique_results

    def _category_from_raw(self, raw_class):
        """
        Определяет категорию объекта по его классу.
        
        Args:
            raw_class: Сырое значение класса объекта
            
        Returns:
            str or None: Категория объекта или None
        """
        if pd.isna(raw_class):
            return None
            
        text = str(raw_class).strip().lower()
        if not text or text == "nan":
            return None
            
        if "армат" in text:
            return "Арматура"
            
        return "Контроль"

    def _category_from_filename(self, file_path):
        """
        Определяет категорию по имени файла.
        
        Args:
            file_path: Путь к файлу
            
        Returns:
            str: Категория объекта
        """
        name = os.path.basename(file_path).lower()
        
        if ("арматур" in name) or ("регулиру" in name):
            return "Арматура"
            
        return "Контроль"

    def _build_result_row(self, value, status, color, category):
        """
        Создает строку результата.
        
        Args:
            value: Значение
            status: Статус совпадения
            color: Цвет для отображения
            category: Категория объекта
            
        Returns:
            dict: Словарь с данными результата
        """
        return {
            "value": value,
            "status": status,
            "color": color,
            "category": category or "Контроль"
        }

    def parse_all_related(self, kks, scheme_file, tables_folder):
        """
        Парсит все связанные файлы для заданного кода KKS.
        
        Args:
            kks: Код KKS для поиска
            scheme_file: Путь к файлу схемы
            tables_folder: Папка с табличными файлами
        """
        # Поиск связанных файлов
        matching_files = self._find_related_files(kks, tables_folder)

        # Создание Excel файла с результатами
        wb = Workbook()
        ws_valves = wb.active
        ws_valves.title = "Арматура"
        ws_control = wb.create_sheet(title="Контроль")

        # Обработка каждой категории
        for sheet_name, files in matching_files.items():
            ws = ws_valves if sheet_name == "Арматура" else ws_control
            
            # Сбор и фильтрация результатов
            collected_results = []
            for matching_file in files:
                collected_results.extend(
                    self.parse_two_tables(scheme_file, matching_file)
                )
            
            filtered_results = self._filter_by_category(
                collected_results, sheet_name
            )
            unique_results = self._deduplicate_results(filtered_results)
            final_results = self._resolve_duplicates(unique_results)
            
            self._write_results_to_sheet(ws, final_results)

        # Очистка лишних строк
        self._cleanup_sheets([ws_valves, ws_control])

        # Сохранение результата
        parent_dir = os.path.dirname(tables_folder)
        output_path = os.path.join(parent_dir, f'{kks}.xlsx')
        wb.save(output_path)
        print(f"✅ Результаты сохранены в '{kks}.xlsx'")
    
    def _find_related_files(self, kks, search_folder):
        """
        Находит все связанные файлы для заданного кода KKS.
        
        Args:
            kks: Код KKS для поиска
            search_folder: Папка для поиска
            
        Returns:
            dict: Словарь с категориями и списками файлов
        """
        matching_files = {"Арматура": [], "Контроль": []}
        
        for filename in os.listdir(search_folder):
            full_path = os.path.join(search_folder, filename)
            
            if not os.path.isfile(full_path) or kks not in filename:
                continue
            
            if ("Запорная арматура" in filename or 
                "Регулирующая арматура" in filename):
                matching_files["Арматура"].append(full_path)
            elif ("Точки контроля" in filename or 
                  "Точки теплотехнического контроля" in filename):
                matching_files["Контроль"].append(full_path)
                
        return matching_files
    
    def _filter_by_category(self, results, category):
        """
        Фильтрует результаты по категории.
        
        Args:
            results: Список результатов
            category: Категория для фильтрации
            
        Returns:
            list: Отфильтрованные результаты
        """
        return [
            row for row in results
            if (category == "Арматура" and row.get("category") == "Арматура")
            or (category == "Контроль" and row.get("category") != "Арматура")
        ]
    
    def _resolve_duplicates(self, results):
        """
        Разрешает дубликаты, оставляя строки со статусом "Совпадает".
        
        Args:
            results: Список результатов
            
        Returns:
            list: Результаты без дубликатов
        """
        value_groups = {}
        
        for row in results:
            value = row["value"]
            if value not in value_groups:
                value_groups[value] = []
            value_groups[value].append(row)
        
        final_results = []
        
        for value, rows in value_groups.items():
            if len(rows) > 1:
                # Приоритет строкам со статусом "Совпадает"
                matching_row = next(
                    (row for row in rows if row.get("status") == "Совпадает"), 
                    None
                )
                if matching_row:
                    final_results.append(matching_row)
                else:
                    final_results.extend(rows)
            else:
                final_results.extend(rows)
                
        return final_results
    
    def _cleanup_sheets(self, worksheets):
        """
        Удаляет служебные строки из листов Excel.
        
        Args:
            worksheets: Список листов для очистки
        """
        for ws in worksheets:
            rows_to_delete = []
            
            for row_idx in range(1, ws.max_row + 1):
                cell_value = ws.cell(row=row_idx, column=1).value
                if cell_value and str(cell_value).strip() == "Краткое наименование":
                    rows_to_delete.append(row_idx)
            
            # Удаление в обратном порядке
            for row_idx in reversed(rows_to_delete):
                ws.delete_rows(row_idx, 1)