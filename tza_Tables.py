import sys

from PyQt5.QtGui import QFont
from PyQt5.QtWidgets import QApplication, QWidget, QVBoxLayout, QPushButton, QFileDialog, QLabel, QMessageBox
import pypdfium2 as pdfium
import pdfplumber
from openpyxl.workbook import Workbook
import os
import re

patterns = [r'Технологическое задание.[а-яА-Я ]+контроля',
            r'Технологическое задание.[а-яА-Я ]+арматура',
            r"Технологическое задание. Регулирующая арматура",
            r"Технологическое задание. Запорная арматура",
            r"Технологическое задание. Дистанционные указатели положения",
            r"Технологическое задание. Регуляторы",
            r"Технологическое задание. Механизмы",
            r"Технологическое задание. Точки контроля",
            r"Технологическое задание. Точки теплотехнического контроля",
            r"Технологическое задание. Точки радиационного контроля"
            r"Технологическое задание. ДУП",
            r"Технологическое задание. Алгоритмы",
            r"Технологическое задание. Сигналы",
            r"Перечень алгоритмов",r"Перечень оборудования", r"Перечень сигналов"]
#patterns = [r"Технологическое задание. Точки контроля"]
combined_pattern = '|'.join(map(re.escape, patterns))

pat_for_imp_pages = r'[A-Z]{3}\.[0-9]{4}\.[0-9]{1,2}[A-Z]{0,3}\.[0-9A-Z]{1,3}\.[A-Z]{2}\.[A-Z]{2}0{3}[0-9]/[0-9]+(?:\.[0-9]+)?'

#----------------------------------------------
class myPdfReader:

    def __init__(self):
        self.names = []
        self.start = 6


    def find_blocks(self, filePath):  # c pdfium ///словарь {страница: соответствующая кодировка справа внизу страницы}
        numb_imp_pages = dict()
        with open(filePath, 'rb') as file:
            data = file.read()
            pdf = pdfium.PdfDocument(data)
            for i in range(self.start, len(pdf)):
                page = pdf.get_page(i)
                textpage = page.get_textpage()
                text = textpage.get_text_bounded()
                #pattern = r'[A-Z]{3}\.[0-9]{4}\.[0-9]{2}[A-Z]{3}\.[A-Z]{3}\.[A-Z]{2}\.[A-Z]{2}0{3}[0-9]/[0-9]+(?:\.[0-9]+)?'
                #pattern = r'[A-Z]{3}\.[0-9]{4}\.[0-9]{2}[A-Z]{3}\.[0-9]\.[A-Z]{2}\.[A-Z]{2}0{3}[0-9]/[0-9]+(?:\.[0-9]+)?'
                match = re.findall(pat_for_imp_pages, text)
                result = match[0].split('/')[-1]
                numb_imp_pages[f"{i+1}"] = float(result)
        return numb_imp_pages

    def find_imp_pages(self, filePath):               #функция для поиска страниц с интересующими таблицами c pdfium
        with open(filePath, 'rb') as file:
            data = file.read()
            pdf = pdfium.PdfDocument(data)
            important_pages = []
            blocks = self.find_blocks(filePath)
            imp_key = 0
            j = -1
            for i in range(self.start, len(pdf) - 1):
                page = pdf.get_page(i)
                textpage = page.get_textpage()
                text = textpage.get_text_bounded()
                matches = re.findall(combined_pattern, text)
                if matches:
                    for important_thing in matches:
                        j += 1
                        important_pages.append({str(i + 1)})
                        self.names.append(important_thing)
                        imp_key = int(blocks[f'{i + 1}'])
                if imp_key == int(blocks[f'{i + 1}']):
                    important_pages[j].add(str(i + 1))
            print(self.names)
        return important_pages

    def extract_imp_tables(self, filePath, pages): #функция для выгрузки таблиц
        pages = sorted(pages)
        my_tables = []
        for page in pages:
            pdf = pdfplumber.open(filePath)
            p = pdf.pages[page-1]
            table = p.extract_table() #ВОТ ТУТ ПРОБЛЕМЫ СО СКРЫТЫМ ТЕКСТОМ
            my_tables.append(table)
        return my_tables


class PDFProcessorApp(QWidget):
    def __init__(self):
        super().__init__()

        # Инициализация интерфейса
        self.initUI()

    def initUI(self):
        self.setWindowTitle("PDF Processor")
        self.setGeometry(400, 400, 700, 300)
        self.setStyleSheet("background-color: rgb(237, 222, 255);")

        # Создание layout
        layout = QVBoxLayout()

        # Метки и кнопки
        self.label_input = QLabel("Выберите папку с PDF файлами", self)
        layout.addWidget(self.label_input)

        self.btn_select_input = QPushButton("Выбрать папку с PDF", self)
        self.btn_select_input.clicked.connect(self.select_input_folder)
        layout.addWidget(self.btn_select_input)

        self.label_output = QLabel("Выберите папку для сохранения обработанных файлов", self)
        layout.addWidget(self.label_output)

        self.btn_select_output = QPushButton("Выбрать папку для сохранения", self)
        self.btn_select_output.clicked.connect(self.select_output_folder)
        layout.addWidget(self.btn_select_output)

        self.btn_process = QPushButton("Обработать PDF файлы", self)
        self.btn_process.clicked.connect(self.process_pdfs)
        layout.addWidget(self.btn_process)

        self.setLayout(layout)

    def select_input_folder(self):
        # Открытие диалогового окна для выбора папки с PDF файлами
        folder = QFileDialog.getExistingDirectory(self, "Выберите папку с PDF файлами")
        if folder:
            self.input_folder = folder
            self.label_input.setText(f"Выбрана папка: {self.input_folder}")

    def select_output_folder(self):
        # Открытие диалогового окна для выбора папки для сохранения обработанных файлов
        folder = QFileDialog.getExistingDirectory(self, "Выберите папку для сохранения обработанных файлов")
        if folder:
            self.output_folder = folder
            self.label_output.setText(f"Выбрана папка: {self.output_folder}")

    def process_pdfs(self):
        if hasattr(self, 'input_folder') and hasattr(self, 'output_folder'):
            # Получаем список PDF файлов в выбранной папке
            pdf_files = [f for f in os.listdir(self.input_folder) if f.lower().endswith('.pdf')]

            if pdf_files:
                for f in pdf_files:
                    filepath = os.path.abspath(f"{self.input_folder}/{f}")
                    k = myPdfReader()
                    imp_page_groups = k.find_imp_pages(filepath)
                    i = 0
                    kks = f.strip('.pdf')
                    for group in imp_page_groups:
                        imp_pages = [int(x) for x in group]
                        l = myPdfReader()
                        tables = l.extract_imp_tables(filepath, pages=imp_pages)
                        ind = 0
                        list_df = []
                        combined_data = []
                        if k.names.count(k.names[i]) > 1:
                            excelPath = f'{self.output_folder}/{kks}_{k.names[i]}_{i}.xlsx'
                        else:
                            excelPath = f'{self.output_folder}/{kks}_{k.names[i]}.xlsx'

                        workbook = Workbook()
                        workbook.active.title = "CommonList"
                        worksheet = workbook['CommonList']

                        if tables:
                            combined_data.extend(tables[0])

                        for table in tables[1:]:
                            if k.names[i] in ["Перечень сигналов", "Перечень алгоритмов"]:
                                combined_data.extend(table[1:])
                            else:
                                combined_data.extend(table[4:])
                        for row in combined_data:
                            worksheet.append(row)

                        worksheet.insert_rows(1)
                        worksheet['A1'] = f"{k.names[i]}"

                        # Проходим по всем ячейкам
                        for row in worksheet.iter_rows():
                            for cell in row:
                                if isinstance(cell.value, str):  # Проверяем, является ли значение ячейки строкой
                                    # Удаляем (cid:13) из текста
                                    cell.value = cell.value.replace('(cid:13)', '')

                        workbook.save(excelPath)
                        print(f'{kks}_{k.names[i]}.xlsx is ready')
                        i += 1

                self.show_message("Успешно!", "PDF файлы обработаны и сохранены.")
            else:
                self.show_message("Ошибка", "В выбранной папке нет PDF файлов.", error=True)
        else:
            self.show_message("Ошибка", "Не выбраны папки для входных и выходных файлов.", error=True)

    def show_message(self, title, message, error=False):
        # Отображение сообщения об успехе или ошибке
        msg = QMessageBox()
        msg.setIcon(QMessageBox.Critical if error else QMessageBox.Information)
        msg.setWindowTitle(title)
        msg.setText(message)
        msg.exec_()


if __name__ == '__main__':
    app = QApplication(sys.argv)
    app_font = QFont("Calibri", 14, QFont.Light)
    app.setFont(app_font)
    window = PDFProcessorApp()
    window.show()
    sys.exit(app.exec_())
