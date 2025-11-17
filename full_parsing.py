import pandas as pd
import re
import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

class mTP:
    def parse_two_tables(self, file1, file2, ws):
        df1 = pd.read_excel(file1)
        df2 = pd.read_excel(file2)
        #df2 = df2.dropna()

        # Получаем значения из первого столбца
        col1 = df1.iloc[:, 0].astype(str).str.strip()
        col2 = df2.iloc[4:, 0].astype(str).str.strip()

        # Создаем множество ключей для поиска
        keys_set = set(col1)

        # Результаты сравнения
        results = []

        # Словарь для отслеживания использованных ключей (чтобы не дублировать)
        used_keys = set()

        # Проверяем каждую строку из второго файла
        for value in col2:
            matched = False
            for key in keys_set:
                # Ищем точное совпадение в начале строки
                if re.match(r'^' + re.escape(key) + r'\b', value):
                    results.append([key, "Совпадает", "green"])
                    used_keys.add(key)
                    matched = True
                    break
            if not matched:
                #results.append([value, f"Только в таблице {file2}", "red"])
                results.append([value, f"Есть в таблице, нет на схеме", "red"])

        # Добавляем оставшиеся непарные элементы из первого файла
        for key in keys_set:
            if key not in used_keys:
                #results.append([key, f"Только на схеме {file1}", "red"])
                results.append([key, f"Есть на схеме, нет в таблице", "red"])

        headers = ["Значение", "Статус"]
        for col_num, data in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=data)

        # Заполняем данными
        for row_num, (value, status, color) in enumerate(results, 2):
            ws.cell(row=row_num, column=1, value=value)
            ws.cell(row=row_num, column=2, value=status)

        # Цвета
        color_map = {
            "green": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),  # зелёный
            "red": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # красный
        }

        # Применяем стили
        for row_num, (_, _, color) in enumerate(results, 2):
            ws.cell(row=row_num, column=1).fill = color_map[color]
            ws.cell(row=row_num, column=2).fill = color_map[color]

        # Автоподгон ширины столбцов
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

    def parse_all_related(self, kks, file1, search_folder):
        # Список для хранения совпадений
        matching_files = []
        # Перебираем все файлы в указанной папке
        for file in os.listdir(search_folder):
            # Полный путь к файлу
            full_path = os.path.join(search_folder, file)
            # Проверяем, является ли это файлом и содержит ли имя нужное слово
            if os.path.isfile(full_path) and kks in file and (("Запорная арматура" in file) or ("Регулирующая арматура" in file) or ("Точки контроля" in file) or ("Точки теплотехнического контроля" in file)):
                matching_files.append(full_path)

        #создаем эксельку
        wb = Workbook()
        i = 0
        for matching_file in matching_files:
            i += 1
            ws = wb.create_sheet(title=f"{i}")
            self.parse_two_tables(file1, matching_file, ws)

        parent_dir = os.path.dirname(search_folder)
        wb.save(f'{parent_dir}/{kks}.xlsx')
        print(f"✅ Результаты сохранены в '{kks}.xlsx'")



if __name__ == "__main__":
    files_path = "C:/Users/User/Desktop/best_new_new/SCHEMES"
    folder = "C:/Users/User/Desktop/best_new_new/TABLES"
    dir_list = os.listdir(files_path)
    for f in dir_list:
        mtp = mTP()
        kks = f.strip('.xlsx')
        mtp.parse_all_related(kks, f"{files_path}/{f}", folder)